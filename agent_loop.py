"""NL2SQL Agent Loop - agent 层驱动代码.

架构: 预处理(量材) → 材料包 → agent 循环(function calling, 预算≤15) → answer
工具: peek_table / peek_values / query_db(意图渲染+门禁) / execute_sql(门禁)
原则: 机器只给数字和执行, 理解和挑选全归 agent; 错误带自纠线索回喂.
"""

import json
import os
import re
import secrets
import sqlite3
import sys
import threading
from datetime import datetime

from openai import OpenAI
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client

from semantic import registry as semantic_registry
from semantic import planner as semantic_planner
from semantic import engine as semantic_engine
from semantic import formulas as semantic_formulas
from semantic import chart as semantic_chart
from semantic import metric_changes

_MDL_CACHE = {}
DATA_ROOT = "data"
BENCH_DIRS = {
    "spider": os.path.join(DATA_ROOT, "spider_data", "database"),
    "bird": os.path.join(DATA_ROOT, "minidev", "MINIDEV", "dev_databases"),
    "ecommerce": os.path.join(DATA_ROOT, "ecommerce"),
}
SKILLS_DIR = os.path.join(DATA_ROOT, "ecommerce", "skills")
MODEL = "deepseek-flash"
BUDGET = 15
SQL_TIMEOUT = 3
RESULT_SAMPLE_ROWS = 20
PEEK_VALUES_LIMIT = 50
MATERIAL_CHARS_THRESHOLD = 20000
WRITE_WORDS = {"INSERT", "UPDATE", "DELETE", "DROP", "CREATE", "ALTER",
               "ATTACH", "DETACH", "REPLACE", "VACUUM", "PRAGMA", "REINDEX"}

SYSTEM_PROMPT = """# Role
You are a data analysis agent.
{database_context}

# Working principles
1. Semantic preservation: keep the question's semantics — set vs multiset (entity lists deduplicated, fact records preserved), grouping grain, entity identity.
2. Data fidelity: report values as stored; no transformation unless requested.
3. Grounding: every literal in a query must come from observed data, never from speculation.
4. Honest delivery: report empty results with verified reasons; keep observing when uncertain; never fabricate.
5. Three states of "no result": legitimate empty set (data exists, truly zero -> answer 0); missing source (table/data not available -> say missing, never a number); query failure (show error, fix within budget). Query beyond data coverage -> say not covered, never 0.
6. Evidence annotation: every number in the final answer must cite its execution reference (Qn) — and ONLY the queries whose data it actually derives from. Qn numbering is session-continuous (never restate "本轮/首轮"). Statements about missing data must cite the observed evidence instead. Format: end the user-facing conclusion section with the marker line "###EVIDENCE###" — before the marker: concise conclusion, each number cited with its (Qn), no markdown bold, no large audit blocks; after the marker: audit evidence (Qn citations, caliber, cross-checks).

# Output format examples

Good answer (correct):
结论区(简洁, 每个数字引用它真实来源的 Qn, 无粗体):
  前7天商品支付金额 1,000,000 元 (Q4)；与近7天 800,000 元 (Q2) 相比减少 200,000 元, 降幅 20.0%
  ###EVIDENCE###
  Q4: SELECT SUM(...) ... → 1,000,000 元（前7天窗口 [08-30, 09-06)）
  Q2: 首轮已查得近7天 800,000 元
  逐日明细: ...(交叉核对)

Bad answer (forbidden — 不要这样写):
  **结论：** 前7天 = **1,000,000 元**（本轮 Q1）……   ← ✗ 粗体、"本轮 Q1"式含糊引用(Qn 会话级连续, 无需本轮/首轮前缀)、把引用当装饰塞满结论
  等等，且缺少 ###EVIDENCE### 分隔符   ← ✗ 无分隔符则证据与结论混在一起

# Full example (few-shot, 从用户问题到最终回答的完整格式)

用户: "近7天比前7天变化多少？"

你的最终回答(完整原文, 一字不差这种格式):
结论: 近7天商品支付金额 800,000 元 (Q2)，较前7天 1,000,000 元 (Q4) 减少 200,000 元, 降幅 20.0%; 支付订单数 4,000→5,000 单同步减少 1,000 单 (Q2, Q4), 单均金额两期持平 (Q2, Q4)。
###EVIDENCE###
Q2: 近7天 [09-06, 09-13) 商品支付金额 SUM(amount)=800,000 元, 支付订单数 4,000 单。
Q4: 前7天 [08-30, 09-06) 商品支付金额 SUM(amount)=1,000,000 元, 支付订单数 5,000 单。
计算: 800,000-1,000,000=-200,000 元; -200,000/1,000,000=-20.0%; 4,000-5,000=-1,000 单。
口径: 商品支付金额=SUM(order_items.amount), 支付时间口径, 仅 status='paid', 不含退款 (技能卡默认口径)。
(要点: 结论区=一句话总结+精确 Qn 引用, 无粗体; 证据区=每个 Qn 的 SQL 事实+计算过程+口径)

# Constraints
Read-only only; finish within the step budget. When you have the answer, reply with the result and a one-sentence Chinese explanation of your choices (no tool call)."""


# ---------------- 预处理层 ----------------

def get_tables(db_path: str) -> list:
    conn = sqlite3.connect(db_path)
    try:
        return [r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
    finally:
        conn.close()


def skill_index(skills_dir: str) -> list:
    """技能索引: 解析每个 .md 的 frontmatter(name/summary)."""
    out = []
    if not os.path.isdir(skills_dir):
        return out
    for fn in sorted(os.listdir(skills_dir)):
        if not fn.endswith(".md"):
            continue
        path = os.path.join(skills_dir, fn)
        head = open(path).read(600)
        m = re.match(r"^---\n(.*?)\n---", head, re.S)
        meta = {}
        if m:
            for line in m.group(1).splitlines():
                if ":" in line:
                    k, _, v = line.partition(":")
                    meta[k.strip()] = v.strip()
        out.append({"name": meta.get("name", fn[:-3]), "summary": meta.get("summary", ""),
                    "path": path})
    return out


def read_skill(skills_dir: str, name: str) -> dict:
    for s in skill_index(skills_dir):
        if s["name"] == name:
            content = open(s["path"]).read()
            body = re.sub(r"^---\n.*?\n---\n", "", content, flags=re.S)  # 去掉 frontmatter
            return {"skill": name, "content": body.strip()}
    return {"error": "unknown_skill", "available": [s["name"] for s in skill_index(skills_dir)]}


def ask_user(question: str, options: list) -> dict:
    """反问工具: 交互模式真问; 批跑(非 tty)模式返回不可用, agent 应改用卡片默认口径."""
    if os.isatty(0):
        print(f"\n[agent 提问] {question}")
        for i, opt in enumerate(options, 1):
            print(f"  {i}. {opt}")
        ans = input("选择编号或直接输入: ").strip()
        if ans.isdigit() and 1 <= int(ans) <= len(options):
            return {"status": "answered", "answer": options[int(ans) - 1]}
        return {"status": "answered", "value": ans or options[0]}
    return {"status": "unavailable",
            "note": "批跑模式无用户可问; 按技能卡片中标注的默认口径执行, 并在答案中声明口径"}


def get_data_cutoff(db_path: str) -> str:
    """数据覆盖截止时间(确定性): 取 pay_time 最大值; 无此表返回 None."""
    try:
        cols = {r[1] for r in sqlite3.connect(db_path).execute('PRAGMA table_info("orders")')}
        if "pay_time" not in cols:
            return None
        row = sqlite3.connect(db_path).execute("SELECT max(pay_time) FROM orders").fetchone()
        return row[0][:10] if row and row[0] else None
    except sqlite3.Error:
        return None


def claim_dir(parent: str, prefix: str) -> str:
    """独占认领新目录: 时间戳+随机尾, 同名当场报错换尾重试 — 同秒并发也不撞."""
    os.makedirs(parent, exist_ok=True)
    while True:
        cand = f"{prefix}-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(2)}"
        try:
            os.mkdir(os.path.join(parent, cand))
            return cand
        except FileExistsError:
            continue


def build_material(db_path: str, db_id: str, skills_dir: str = None) -> str:
    """量材 + 组装: 小库给全部列名, 大库只给表名. 加技能索引与数据覆盖声明."""
    tables = get_tables(db_path)
    mdl = os.path.join(os.path.dirname(db_path), "mdl.yaml")
    reg = load_mdl(mdl, db_path=db_path)
    hidden_cols = {
        column
        for model_name in (reg.models if reg else {})
        for column in reg.models[model_name].get("columns", [])
        if column.get("hidden", False)
    }
    conn = sqlite3.connect(db_path)
    try:
        lines = []
        for t in tables:
            cols = [r[1] for r in conn.execute(f'PRAGMA table_info("{t}")')
                    if r[1] not in hidden_cols]
            lines.append(f"- {t}({', '.join(cols)})")
        text = "\n".join(lines)
    finally:
        conn.close()
    if len(text) <= MATERIAL_CHARS_THRESHOLD:
        material = "数据库目录与全部列名:\n" + text
    else:
        material = "数据库目录(大库, 仅表名):\n" + "\n".join(f"- {t}" for t in tables)
    cutoff = get_data_cutoff(db_path)
    if cutoff:
        material += (f"\n\n数据覆盖: 截至 {cutoff}。查询范围超出该日期时, "
                     "如实回答'数据未覆盖', 不得回答 0。")
    skills = skill_index(skills_dir) if skills_dir else []
    if skills:
        idx = "\n".join(f"- [{s['name']}] {s['summary']}（用 read_skill(\"{s['name']}\") 读取全文）"
                        for s in skills)
        material += f"\n\n可用技能(处理对应话题前先读):\n{idx}"
    if reg is not None:
        material += "\n\n" + semantic_registry.describe(reg)
        material += ("\n\n查询上述标准指标(含平台指标如 访客数)请用 query_metrics 点菜单"
                     "(metrics/dimensions/time_range), 不要自己写聚合 SQL、"
                     "也不要直接调用已登记的平台 MCP 工具(get_shop_traffic 等——"
                     "它们的取数已由语义层接管, 直接调用会被拦截)。"
                     "探索/明细/复杂查询再用 execute_sql 或 peek_*。"
                     "需要两期对比时，在 query_metrics 里加 compare:{start,end}，"
                     "直接得到差额/降幅，不要自己查两次相减。"
                     "要图时用 chart(基于 query_metrics 的 spec)。"
                     "需要修改口径时只能用 propose_metric_change 提议，不能批准或直接改 mdl.yaml。"
                     "query_metrics 结果中 null = 不适用(分母为0/数据缺失), "
                     "如实转述'不适用', 禁止写 0 或编数。")
        pending = metric_changes.list_pending(os.path.dirname(db_path))
        if pending:
            grouped = {}
            for draft in pending:
                grouped.setdefault(draft.get("metric"), []).append(draft.get("id"))
            notices = "；".join(
                f"{metric} 口径有 {len(ids)} 条变更待审批(id={','.join(ids)})"
                for metric, ids in grouped.items()
            )
            material += "\n\n待审批口径变更: " + notices
    return material


def load_mdl(path: str, db_path: str = None):
    """加载缓存语义层(含声明主键的真实唯一性校验); 无文件返回 None。

    以 mtime 作为缓存版本，审批脚本改写 mdl.yaml 后同一进程也会重载，
    避免 agent 继续使用旧口径。
    """
    if not os.path.exists(path):
        return None
    mtime = os.stat(path).st_mtime_ns
    cached = _MDL_CACHE.get(path)
    if cached and cached[0] == mtime:
        return cached[1]
    reg = semantic_registry.load(path, db_path=db_path)
    _MDL_CACHE[path] = (mtime, reg)
    return reg


def run_query_metrics(spec: dict, db_path: str, mcp_bridge=None, tenant=1) -> dict:
    """QuerySpec -> 规划器 -> 执行. 规划/校验错误以结构化 payload 回喂 agent."""
    mdl = os.path.join(os.path.dirname(db_path), "mdl.yaml")
    reg = load_mdl(mdl, db_path=db_path)
    if reg is None:
        return {"error": "no_semantic_layer", "reason": f"未找到 {mdl}"}
    try:
        plan = semantic_planner.Planner(reg).plan(spec, tenant=tenant)
        out = semantic_engine.run(plan, db_path, mcp_bridge=mcp_bridge)
        out["sql"] = " ; ".join(out.get("node_sql", []))
        return out
    except semantic_planner.PlanError as e:
        return e.payload
    except semantic_registry.RegistryError as e:
        return {"error": "registry_error", "reason": str(e)}


def chart(spec: dict, chart_type: str, x: str, y: str, db_path: str,
          mcp_bridge=None, tenant=1, _return_query: bool = False) -> dict:
    """查询标准指标并落盘 ECharts option；``_return_query`` 仅供 dispatch 记 Qn。"""
    if chart_type not in semantic_chart.chart_types():
        return {"error": "unknown_chart_type"}
    query = run_query_metrics(spec, db_path, mcp_bridge=mcp_bridge, tenant=tenant)
    if query.get("error"):
        return query
    if query.get("status") not in ("ok", "empty"):
        result = {"error": "query_failed", "reason": query.get("status")}
        if _return_query:
            result["_query_result"] = query
        return result
    columns = query.get("columns") or []
    if x not in columns:
        result = {"error": "unknown_chart_column", "column": x,
                  "available": columns}
        if _return_query:
            result["_query_result"] = query
        return result
    if y not in columns:
        result = {"error": "unknown_chart_column", "column": y,
                  "available": columns}
        if _return_query:
            result["_query_result"] = query
        return result
    try:
        option = semantic_chart.to_echarts(columns, query.get("rows") or [],
                                           chart_type, x, y)
    except ValueError as exc:
        code = str(exc)
        if code == "unknown_chart_type":
            result = {"error": code}
        else:
            result = {"error": "unknown_chart_column", "reason": code}
        if _return_query:
            result["_query_result"] = query
        return result

    os.makedirs("reports", exist_ok=True)
    chart_id = f"c-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(2)}"
    path = os.path.join("reports", chart_id + ".chart.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(option, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    result = {"status": "saved", "path": path, "option": option,
              "n_rows": query.get("n_rows", len(query.get("rows") or [])),
              "columns": columns}
    if _return_query:
        result["_query_result"] = query
    return result


# ---------------- 工具实现（确定性） ----------------

def peek_table(db_path: str, table: str, columns: list = None) -> dict:
    tables = get_tables(db_path)
    if table not in tables:
        return {"error": "unknown_table", "available_tables": tables}
    conn = sqlite3.connect(db_path)
    try:
        all_cols = [r[1] for r in conn.execute(f'PRAGMA table_info("{table}")')]
        wanted = [c for c in all_cols if not columns or c in columns]
        fks = []
        pk = [r[1] for r in conn.execute(f'PRAGMA table_info("{table}")') if r[5]]
        if pk:
            fks.append(f"{table}.主键 = {', '.join(pk)}")
        seen = set()
        for r in conn.execute(f'PRAGMA foreign_key_list("{table}")'):
            entry = f"{table}.{r[3]} -> {r[2]}.{r[4]}"
            if entry not in seen:
                seen.add(entry)
                fks.append(entry)
        stats = []
        total = conn.execute(f'SELECT count(*) FROM "{table}"').fetchone()[0]
        for c in wanted:
            d = conn.execute(f'SELECT count(DISTINCT "{c}") FROM "{table}"').fetchone()[0]
            nn = conn.execute(f'SELECT count("{c}") FROM "{table}"').fetchone()[0]
            avg = conn.execute(
                f'SELECT avg(length(CAST("{c}" AS TEXT))) FROM "{table}" WHERE "{c}" IS NOT NULL'
            ).fetchone()[0] or 0
            stats.append(f"  {c}: 不同值{d}, 非空率{round(nn*100/total,1) if total else 0}%, 平均长度{int(avg)}")
        return {"table": table, "外键": fks, "列统计": stats}
    finally:
        conn.close()


def peek_values(db_path: str, table: str, column: str, keyword: str = None) -> dict:
    cols = [r[1] for r in sqlite3.connect(db_path).execute(f'PRAGMA table_info("{table}")')]
    if column not in cols:
        return {"error": "unknown_column", "available_columns": cols}
    conn = sqlite3.connect(db_path)
    try:
        if keyword:
            rows = conn.execute(
                f'SELECT DISTINCT CAST("{column}" AS TEXT) FROM "{table}" '
                f'WHERE "{column}" IS NOT NULL AND CAST("{column}" AS TEXT) LIKE ? LIMIT ?',
                (f"%{keyword}%", PEEK_VALUES_LIMIT)).fetchall()
        else:
            rows = conn.execute(
                f'SELECT DISTINCT CAST("{column}" AS TEXT) FROM "{table}" '
                f'WHERE "{column}" IS NOT NULL LIMIT ?', (PEEK_VALUES_LIMIT,)).fetchall()
        d = conn.execute(f'SELECT count(DISTINCT "{column}") FROM "{table}"').fetchone()[0]
        out = {"table": table, "column": column, "total_distinct": d,
               "values": [r[0][:80] for r in rows]}
        if d > len(rows):
            out["note"] = f"共 {d} 种值, 此为前 {len(rows)} 种; 用 keyword 精确定位"
        return out
    finally:
        conn.close()


def gate_check(sql: str, db_path: str, bench: str, db_id: str) -> dict:
    body = sql.strip().removeprefix("```sql").removesuffix("```").strip().rstrip(";")
    if not body:
        return {"error": "sql_rejected", "reason": "空 SQL"}
    head = body.split(None, 1)[0].upper()
    if head != "SELECT":
        return {"error": "sql_rejected", "reason": f"只允许 SELECT, 收到: {head}"}
    if ";" in body:
        return {"error": "sql_rejected", "reason": "禁止多语句"}
    upper = body.upper()
    for w in WRITE_WORDS:
        if w in upper.split() or f"{w} " in upper:
            return {"error": "sql_rejected", "reason": f"禁止写操作关键词: {w}"}
    vdb = os.path.join(DATA_ROOT, "profiles", f"{bench}_{db_id}_values.db")
    if os.path.exists(vdb):
        conn = sqlite3.connect(vdb)
        try:
            literals = [l for l in set(re.findall(r"'([^']*)'", body)) if l and not l.isdigit()]
            for lit in literals:
                if len(lit) < 3:
                    continue
                rows = conn.execute("SELECT value FROM vi WHERE value_lower=? LIMIT 1",
                                    (lit.lower(),)).fetchall()
                if not rows:
                    return {"error": "value_not_found", "value": lit,
                            "reason": "该值不在值索引中, 请先 peek_values 验证"}
        finally:
            conn.close()
    else:
        conn = sqlite3.connect(db_path)
        conn.execute("PRAGMA query_only=ON")
        for lit in set(re.findall(r"'([^']*)'", body)):
            if lit and not lit.isdigit() and len(lit) >= 3:
                found = conn.execute(
                    f'SELECT 1 FROM ({body}) LIMIT 0') is not None  # 无索引时跳过字面值预检
        conn.close()
    return {"ok": True, "sql": body}


def run_sql(sql: str, db_path: str) -> dict:
    body = sql.strip().removeprefix("```sql").removesuffix("```").strip().rstrip(";")
    conn = sqlite3.connect(db_path, timeout=SQL_TIMEOUT)
    conn.execute("PRAGMA query_only=ON")
    result = []
    timer = threading.Timer(SQL_TIMEOUT, lambda: result.append(("timeout", None, None)))
    timer.start()
    try:
        cur = conn.execute(body)
        result.append(("ok", cur.fetchall(), [d[0] for d in (cur.description or [])]))
    except sqlite3.Error as e:
        result.append(("error", str(e), None))
    finally:
        timer.cancel()
        conn.close()
    kind, payload, cols = result[0]
    if kind == "timeout":
        return {"error": "timeout", "reason": f"超过 {SQL_TIMEOUT}s"}
    if kind == "error":
        return {"error": "sql_error", "reason": payload}
    out = {"status": "empty" if not payload else "ok", "n_rows": len(payload),
           "columns": cols or [], "rows_all": payload or []}
    if payload:
        out["rows"] = payload if len(payload) <= RESULT_SAMPLE_ROWS else payload[:5]
        if len(payload) > RESULT_SAMPLE_ROWS:
            out["note"] = f"共 {len(payload)} 行, 此为前 5 行"
    return out


def render_intent(intent: dict, db_path: str) -> dict:
    """六槽位 → SQL. 列校验确定性执行; 单表 JOIN 由 agent 自理(放意图 schema 外)."""
    table = intent.get("table")
    if not table:
        return {"error": "invalid_intent", "reason": "缺少 table"}
    real_cols = {r[1] for r in sqlite3.connect(db_path).execute(f'PRAGMA table_info("{table}")')}
    if not real_cols:
        return {"error": "unknown_table", "available_tables": get_tables(db_path)}
    metric = intent.get("metric")
    filters = intent.get("filters") or []
    if not metric and not filters:
        return {"error": "invalid_intent", "reason": "metric 与 filters 至少其一"}
    cols_used = []
    if metric and metric.get("column"):
        cols_used.append(metric["column"])
    cols_used += [f["column"] for f in filters]
    cols_used += (intent.get("group_by") or [])
    ob = intent.get("order_by") or {}
    if ob.get("by"):
        cols_used.append(ob["by"])
    unknown = [c for c in cols_used if c not in real_cols]
    if unknown:
        return {"error": "unknown_column", "unknown": unknown, "available_columns": sorted(real_cols)}
    where = []
    for f in filters:
        col, op, val = f["column"], (f.get("op") or "=").upper(), f.get("value")
        if op == "=":
            where.append(f"{col} = '{val}' COLLATE NOCASE" if isinstance(val, str) else f"{col} = {val}")
        elif op in (">", "<"):
            where.append(f"{col} {op} {val}")
        elif op == "like":
            where.append(f"{col} LIKE '%{val}%'")
        elif op == "in":
            items = ", ".join(f"'{v}'" if isinstance(v, str) else str(v) for v in (val or []))
            where.append(f"{col} IN ({items})")
    select = "COUNT(*)"
    if metric:
        fn, col = metric.get("fn", "count"), metric.get("column")
        if fn not in ("count", "sum", "avg", "max", "min"):
            return {"error": "invalid_intent", "reason": f"未知聚合: {fn}"}
        select = f"{fn.upper()}({'*' if not col else col})"
    gb = intent.get("group_by") or []
    if gb:
        select = f"{', '.join(gb)}, {select}" if metric and metric.get("column") else f"{', '.join(gb)}, {select}" if metric else ", ".join(gb)
    sql = f"SELECT {select} FROM {table}"
    if where:
        sql += " WHERE " + " AND ".join(where)
    if gb:
        sql += " GROUP BY " + ", ".join(gb)
    if ob.get("by"):
        sql += f" ORDER BY {ob['by']} {ob.get('dir', 'ASC').upper()}"
    if intent.get("limit"):
        sql += f" LIMIT {int(intent['limit'])}"
    return {"ok": True, "sql": sql}


# ---------------- 计算器 + 报告（确定性） ----------------

FORMULA_LIST = {
    "delta":        ("差额 = 本期 - 基期", ["curr", "prev"]),
    "pct_change":   ("降幅% = (本期 - 基期) / 基期 * 100", ["curr", "prev"]),
    "share":        ("占比% = 部分 / 整体 * 100", ["part", "total"]),
    "contribution": ("净贡献% = 部分变化 / 净变化 * 100", ["part_delta", "net_delta"]),
    "avg":          ("平均/比值 = 总量 / 数量", ["total", "count"]),
    "closure":      ("闭合差 = 各部分之和 - 总量(应为0)", ["parts", "total"]),
    "lmdi":         ("LMDI乘法分解 = 各因素对数均值贡献; 入参 total_curr/total_prev/factors",
                     ["total_curr", "total_prev", "factors"]),
}


def _lmdi_tolerance(args: dict, db_path: str = None) -> float:
    """从当前 MDL 的对账声明取 LMDI 容差；没有匹配声明时使用 2%。"""
    default = 0.02
    mdl_path = (os.path.join(os.path.dirname(db_path), "mdl.yaml")
                if db_path else os.path.join(DATA_ROOT, "ecommerce", "mdl.yaml"))
    try:
        reg = load_mdl(mdl_path, db_path=db_path)
    except (OSError, sqlite3.Error, semantic_registry.RegistryError, TypeError, ValueError):
        return default
    if reg is None:
        return default

    factors = args.get("factors") or {}
    rules = reg.reconciliations
    if isinstance(rules, dict):
        rules = list(rules.values())
    # lmdi 入参没有总量指标名，因此优先按因素名集合匹配；同名集合
    # 的声明保留其 YAML 顺序，找不到精确匹配时再采用首条规则。
    for rule in rules:
        if set(rule.get("factors") or []) == set(factors):
            value = rule.get("tolerance", default)
            if isinstance(value, (int, float)) and value >= 0:
                return float(value)
    for rule in rules:
        value = rule.get("tolerance", default)
        if isinstance(value, (int, float)) and value >= 0:
            return float(value)
    return default


def calculate(formula: str, args: dict, db_path: str = None) -> dict:
    """具名公式计算器: 确定性; 返回高精度 value 与 2 位 display; 边界返回不适用, 绝不返回 inf."""
    if formula not in FORMULA_LIST:
        return {"error": "unknown_formula", "available_formulas": list(FORMULA_LIST)}
    try:
        if formula == "delta":
            v = args["curr"] - args["prev"]
        elif formula == "pct_change":
            if args["prev"] == 0:
                return {"status": "undefined", "formula": formula,
                        "note": "基期为0, 降幅不适用(禁止无穷大)"}
            v = (args["curr"] - args["prev"]) / args["prev"] * 100
        elif formula == "share":
            if args["total"] == 0:
                return {"status": "undefined", "formula": formula, "note": "整体为0, 占比不适用"}
            v = args["part"] / args["total"] * 100
        elif formula == "contribution":
            if args["net_delta"] == 0:
                return {"status": "undefined", "formula": formula,
                        "note": "净变化为0, 净贡献比例不稳定, 只报金额与方向"}
            v = args["part_delta"] / args["net_delta"] * 100
        elif formula == "avg":
            if args["count"] == 0:
                return {"status": "undefined", "formula": formula, "note": "数量为0, 平均不适用"}
            v = args["total"] / args["count"]
        elif formula == "lmdi":
            v = semantic_formulas.lmdi(
                args["total_curr"], args["total_prev"], args["factors"],
                tolerance=_lmdi_tolerance(args, db_path=db_path))
        else:  # closure
            v = sum(args["parts"]) - args["total"]
    except (KeyError, TypeError, ValueError, ZeroDivisionError) as e:
        return {"error": "bad_args", "reason": str(e), "expected_args": FORMULA_LIST[formula][1]}
    if formula == "lmdi":
        return {"status": "ok" if v.get("status") == "ok" else v.get("status"),
                "formula": formula, "formula_text": FORMULA_LIST[formula][0],
                "inputs": args, **v}
    return {"status": "ok", "formula": formula, "formula_text": FORMULA_LIST[formula][0],
            "inputs": args, "value": v, "display": round(v, 2)}


def _collect_numbers(obj, out: set):
    """递归收集任意结构里的所有数字(用于对账)."""
    if isinstance(obj, bool) or obj is None:
        return
    if isinstance(obj, (int, float)):
        out.add(float(obj))
    elif isinstance(obj, str):
        for m in re.findall(r"-?\d[\d,]*(?:\.\d+)?", obj):
            try:
                out.add(float(m.replace(",", "")))
            except ValueError:
                pass
    elif isinstance(obj, dict):
        for v in obj.values():
            _collect_numbers(v, out)
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            _collect_numbers(v, out)


def _ledger_values(entry: dict) -> set:
    vals = set()
    _collect_numbers(entry.get("rows"), vals)
    _collect_numbers(entry.get("payload"), vals)
    return vals


def check_report(text: str, executions: dict) -> list:
    """报告正文对账: 每个数字必须能在它所引用的 Qn 的账本记录里找到(容差 0.02).

    扫描前先剔除"非数据数字": 日期、引用记号(Qn)、行首序号、报告ID/路径、时间计数(近7天).
    """
    issues = []
    body = text.split("###EVIDENCE###", 1)[0]
    body = DATE_LIKE.sub(" ", body)
    body = re.sub(r"Q\d+", " ", body)                              # 引用记号本身
    body = re.sub(r"[\w./\\-]*r-\d{8}-\d{6}[\w./\\-]*", " ", body)  # 报告ID/文件名/路径
    body = re.sub(r"(?m)^\s*\d+\s*[\.、)]\s*", " ", body)          # 行首列表序号
    body = re.sub(r"\d+(?=\s*[天周月年日次])", " ", body)           # 近7天/前7天 这类时间计数
    cite_re = re.compile(r"[（(]\s*(Q\d+(?:\s*[,，]\s*Q\d+)*)\s*[)）]")
    num_re = re.compile(r"-?\d[\d,]*(?:\.\d+)?%?")
    last = 0
    for m in cite_re.finditer(body):
        span = body[last:m.start()]
        last = m.end()
        refs = [f"Q{r}" for r in re.findall(r"Q(\d+)", m.group(1))]
        allowed, missing = set(), []
        for r in refs:
            if r not in executions:
                missing.append(r)
            else:
                allowed |= _ledger_values(executions[r])
        if missing:
            issues.append(f"引用的 {', '.join(missing)} 在账本中不存在")
        for tok in num_re.findall(span):
            val = _norm_num(tok)
            if val is None:
                continue
            if not any(abs(val - a) <= 0.02 for a in allowed):
                issues.append(f"数字 {tok} 不在所引 {', '.join(refs)} 的账本记录中")
    return issues


def _report_caliber_section(content: str, state: dict, db_path: str = None) -> str:
    """从报告正文实际引用的 query_metrics Qn 生成口径节，不信任手写指标名。"""
    body = content.split("###EVIDENCE###", 1)[0]
    refs = []
    for number in re.findall(r"\bQ(\d+)\b", body):
        ref = f"Q{number}"
        if ref not in refs:
            refs.append(ref)
    metric_names = []
    for ref in refs:
        entry = state.get("executions", {}).get(ref) or {}
        if entry.get("tool") != "query_metrics":
            continue
        for name in (entry.get("args") or {}).get("metrics") or []:
            if name not in metric_names:
                metric_names.append(name)
    if not metric_names:
        return ""
    if db_path is None:
        db_path = os.path.join(DATA_ROOT, "ecommerce", "ecommerce.sqlite")
    mdl = os.path.join(os.path.dirname(db_path), "mdl.yaml")
    reg = load_mdl(mdl, db_path=db_path)
    if reg is None:
        return ""
    lines = ["## 口径说明（由语义层生成）"]
    for name in metric_names:
        definition = reg.lookup_metric(name)
        if not definition:
            continue
        lines.append(
            f"- {name}：{definition.get('description', '')}；"
            f"单位：{definition.get('unit', '')}；"
            f"负责人：{definition.get('owner', '未登记')}；"
            f"版本：v{definition.get('version', '?')}；"
            f"生效时间：{definition.get('effective_date', '未登记')}"
        )
    return "\n".join(lines) if len(lines) > 1 else ""


def _write_docx(path: str, title: str, body: str, meta: dict,
                caliber_section: str = ""):
    from docx import Document
    doc = Document()
    doc.add_heading(title or "经营分析报告", level=0)
    for line in body.splitlines():
        s = line.rstrip()
        if not s:
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        else:
            doc.add_paragraph(s)
    if caliber_section:
        doc.add_paragraph("")
        for line in caliber_section.splitlines():
            s = line.rstrip()
            if s.startswith("## "):
                doc.add_heading(s[3:], level=2)
            elif s.startswith("- "):
                doc.add_paragraph(s[2:], style="List Bullet")
            elif s:
                doc.add_paragraph(s)
    doc.add_paragraph("")
    doc.add_paragraph(f"数据截止时间: {meta.get('cutoff') or '未取得'}")
    doc.add_paragraph(f"报告ID: {meta['report_id']}  会话: {meta['session']}  生成时间: {meta['created_at']}")
    doc.save(path)


def save_report(title: str, content: str, state: dict) -> dict:
    """报告落盘: 先对账本校验 → 通过才转 Word + 写元数据; 不过则返回 issues 供 agent 修正."""
    if not content.strip():
        return {"error": "empty_report", "reason": "报告正文为空"}
    issues = check_report(content, state.get("executions", {}))
    if issues:
        return {"status": "rejected", "issues": issues,
                "note": "报告未通过对账校验, 请按 issues 修正后重新保存"}
    rid = f"r-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(2)}"
    os.makedirs("reports", exist_ok=True)
    created = datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = {"report_id": rid, "title": title, "session": state.get("session_id"),
            "cutoff": state.get("data_cutoff"), "created_at": created,
            "executions": len(state.get("executions", {}))}
    body = content.split("###EVIDENCE###", 1)[0].strip()
    caliber_section = _report_caliber_section(content, state, state.get("db_path"))
    path = os.path.join("reports", rid + ".docx")
    _write_docx(path, title, body, meta, caliber_section=caliber_section)
    json.dump(meta, open(os.path.join("reports", rid + ".json"), "w"),
              ensure_ascii=False, indent=1)
    return {"status": "saved", "report_id": rid, "path": path,
            "cutoff": meta["cutoff"],
            "caliber_metrics": [
                name for name in re.findall(r"^- ([^：]+)：", caliber_section, re.M)
            ],
            "note": "报告已落盘为 Word; 最终答复请给用户报告全文"}


def _norm_num(tok: str):
    t = tok.strip().replace(",", "").rstrip("%")
    try:
        return float(t)
    except ValueError:
        return None


# ---------------- agent 循环 ----------------

class Trace:
    def __init__(self, path: str):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.fh = open(path, "a")

    def log(self, action: str, **kw):
        self.fh.write(json.dumps(
            {"ts": datetime.now().isoformat(timespec="milliseconds"), "action": action, **kw},
            ensure_ascii=False, default=str) + "\n")
        self.fh.flush()


def tool_schema() -> list:
    return [
        {"type": "function", "function": {
            "name": "peek_table",
            "description": "翻开一张表: 外键接缝 + 每列一行统计(不同值数/非空率/平均长度). 看结构用",
            "parameters": {"type": "object", "properties": {
                "table": {"type": "string"},
                "columns": {"type": "array", "items": {"type": "string"},
                            "description": "可选: 只看指定列"}}}}},
        {"type": "function", "function": {
            "name": "peek_values",
            "description": "查某列实际取值(≤50个). 写过滤条件前验证字面值用",
            "parameters": {"type": "object", "properties": {
                "table": {"type": "string"}, "column": {"type": "string"},
                "keyword": {"type": "string", "description": "可选: 大量取值中筛选"}}}}},
        {"type": "function", "function": {
            "name": "query_db",
            "description": "六槽位意图: 渲染为SQL→校验→执行. 单表查询首选",
            "parameters": {"type": "object", "properties": {
                "table": {"type": "string"},
                "metric": {"type": "object", "nullable": True, "properties": {
                    "fn": {"type": "string", "enum": ["count", "sum", "avg", "max", "min"]},
                    "column": {"type": "string", "nullable": True}}},
                "filters": {"type": "array", "nullable": True, "items": {
                    "type": "object", "properties": {
                        "column": {"type": "string"},
                        "op": {"type": "string", "enum": ["=", ">", "<", "like", "in"]},
                        "value": {}}}},
                "group_by": {"type": "array", "nullable": True, "items": {"type": "string"}},
                "order_by": {"type": "object", "nullable": True, "properties": {
                    "by": {"type": "string"}, "dir": {"type": "string", "enum": ["asc", "desc"]}}},
                "limit": {"type": "integer", "nullable": True},
                "distinct": {"type": "boolean", "nullable": True}}}}},
        {"type": "function", "function": {
            "name": "find_value",
            "description": "反查一个值住在库里的哪些表哪些列(定位器). 无索引时提示改用 peek_values",
            "parameters": {"type": "object", "properties": {
                "value": {"type": "string", "description": "要定位的字面值, 如人名/状态名"}}}}},
        {"type": "function", "function": {
            "name": "read_skill",
            "description": "读取一个技能卡片的全文(指标定义/公式/易错点). 处理对应话题前先读",
            "parameters": {"type": "object", "properties": {
                "name": {"type": "string", "description": "技能名称, 见材料包中的可用技能列表"}}}}},
        {"type": "function", "function": {
            "name": "ask_user",
            "description": "歧义时向用户反问. 有用户时返回其选择; 批跑模式返回不可用(此时应按卡片默认口径执行并在答案中声明)",
            "parameters": {"type": "object", "properties": {
                "question": {"type": "string", "description": "要问的问题"},
                "options": {"type": "array", "items": {"type": "string"},
                            "description": "候选答案列表, 如两种口径"}}}}},
        {"type": "function", "function": {
            "name": "execute_sql",
            "description": "执行一条只读SQL(多表JOIN等意图无法表达时用)",
            "parameters": {"type": "object", "properties": {
                "sql": {"type": "string"}}}}},
        {"type": "function", "function": {
            "name": "query_metrics",
            "description": "按语义层标准指标查询(点菜单): 填指标/维度/时间窗, 机器展开口径。查询已登记指标(GMV/支付订单数/单均金额)必须用它; 探索/明细/复杂查询用 execute_sql",
            "parameters": {"type": "object", "properties": {
                "metrics": {"type": "array", "items": {"type": "string"},
                            "description": "指标名, 见材料包'可用指标'"},
                "dimensions": {"type": "array", "items": {"type": "string"},
                               "description": "分组维度, 如 category"},
                "time_range": {"type": "object", "properties": {
                    "start": {"type": "string"}, "end": {"type": "string"}},
                    "description": "半开区间 [start,end), 绝对日期 YYYY-MM-DD"},
                "compare": {"type": "object", "properties": {
                    "start": {"type": "string"}, "end": {"type": "string"}},
                    "description": "可选: 基期时间窗(半开区间), 与 time_range 对比。给出后每个指标返回 curr/prev/delta/pct"},
                "time_grain": {"type": "string", "enum": ["total", "day"]},
                "filters": {"type": "array", "items": {"type": "object", "properties": {
                    "dimension": {"type": "string"}, "op": {"type": "string"}, "value": {}}}},
                "order_by": {"type": "array", "items": {"type": "object", "properties": {
                    "field": {"type": "string"}, "dir": {"type": "string", "enum": ["asc", "desc"]}}}},
                "limit": {"type": "integer"}},
                "required": ["metrics", "time_range"]}}},
        {"type": "function", "function": {
            "name": "chart",
            "description": "基于 query_metrics 的结果生成 ECharts 图表配置并落盘。支持 bar/line/pie；数字仍由内部 query_metrics 取数并记入 Qn",
            "parameters": {"type": "object", "properties": {
                "spec": {"type": "object", "description": "query_metrics 的 QuerySpec"},
                "chart_type": {"type": "string", "enum": list(semantic_chart.chart_types())},
                "x": {"type": "string", "description": "横轴/饼图名称列"},
                "y": {"type": "string", "description": "数值列(指标名)"}},
                "required": ["spec", "chart_type", "x", "y"]}}},
        {"type": "function", "function": {
            "name": "propose_metric_change",
            "description": "提议修改指标口径，只写待审批草稿，不会批准或修改 mdl.yaml",
            "parameters": {"type": "object", "properties": {
                "metric": {"type": "string"},
                "field": {"type": "string", "description": "可改字段，如 description/unit/expr"},
                "new_value": {},
                "reason": {"type": "string"}},
                "required": ["metric", "field", "new_value", "reason"]}}},
        {"type": "function", "function": {
            "name": "calculate",
            "description": "具名公式计算器(禁止心算). 派生数(差额/降幅/占比/净贡献/平均/闭合/LMDI乘法归因)必须用它算, 结果会进账本并获 Qn",
            "parameters": {"type": "object", "properties": {
                "formula": {"type": "string",
                            "enum": list(FORMULA_LIST)},
                "args": {"type": "object",
                         "description": "公式入参, 键名见公式: delta/pct_change{curr,prev}, share{part,total}, contribution{part_delta,net_delta}, avg{total,count}, closure{parts:[...],total}, lmdi{total_curr,total_prev,factors:{name:{curr,prev}}}"}}}}},
        {"type": "function", "function": {
            "name": "save_report",
            "description": "把报告正文落盘为 Word 并做对账校验(每个数字须能在所引 Qn 的账本记录里找到). 通过返回 saved+路径; 不过返回 rejected+issues, 按 issues 改后重调",
            "parameters": {"type": "object", "properties": {
                "title": {"type": "string", "description": "报告标题"},
                "content": {"type": "string", "description": "报告正文(四段: 问题/数据及结果/原因/改进), 数字后带 (Qn) 引用"}}}}},
    ]


def find_value(value: str, bench: str, db_id: str) -> dict:
    """值索引反查: 这个值住在哪些表哪些列. 无索引时优雅返回."""
    vdb = os.path.join(DATA_ROOT, "profiles", f"{bench}_{db_id}_values.db")
    if not os.path.exists(vdb):
        return {"status": "no_index", "note": "该库未建值索引, 请用 peek_values 按列验证"}
    conn = sqlite3.connect(vdb)
    try:
        rows = conn.execute("SELECT value FROM vi WHERE value_lower=? LIMIT 5", (value.lower(),)).fetchall()
        if rows:
            fuzzy = []
            cand = {r[0] for r in conn.execute(
                "SELECT value FROM vfts WHERE vfts MATCH ? LIMIT 20",
                (f'"{value[:3]}"' if len(value) >= 3 else value,))}
            return {"status": "hit", "value": rows[0][0], "nearby": sorted(cand)[:10]}
        return {"status": "no_match", "note": "所有已索引列中无此值; 可能是语义概念而非数据值, 读目录/画像找"}
    finally:
        conn.close()


# ---------------- MCP 桥接(外部平台数据服务) ----------------

class McpBridge:
    """连接外部 MCP server(独立进程), 把它的工具注册进 agent 工具箱."""

    def __init__(self, command: str, args: list, cwd: str = "."):
        self._params = StdioServerParameters(command=command, args=args, cwd=cwd)
        self.tools = []
        self._loop = None

    def start(self) -> list:
        import asyncio
        import threading
        self._loop = asyncio.new_event_loop()
        self._ready = threading.Event()
        t = threading.Thread(target=self._run, daemon=True)
        t.start()
        self._ready.wait(timeout=30)
        return self.tools

    def _run(self):
        import asyncio
        asyncio.set_event_loop(self._loop)
        self._loop.run_until_complete(self._connect())
        self._loop.run_forever()

    async def _connect(self):
        from mcp import ClientSession
        from mcp.client.stdio import stdio_client
        self._client_cm = stdio_client(self._params)
        read, write = await self._client_cm.__aenter__()
        self._session_cm = ClientSession(read, write)
        self._session = await self._session_cm.__aenter__()
        await self._session.initialize()
        listing = await self._session.list_tools()
        self.tools = [{"name": t.name, "description": t.description or "",
                       "parameters": t.inputSchema} for t in listing.tools]

    def call(self, name: str, args: dict) -> dict:
        import asyncio
        fut = asyncio.run_coroutine_threadsafe(
            self._session.call_tool(name, args), self._loop)
        result = fut.result(timeout=60)
        return json.loads(result.content[0].text)


def dispatch(name: str, args: dict, db_path: str, bench: str, db_id: str,
             trace: Trace, turn: int, state: dict) -> dict:
    trace.log("tool_call", turn=turn, name=name, args=args)
    mcp_names = state.get("mcp_tools", {})
    internal_query = None
    if name == "peek_table":
        result = peek_table(db_path, args["table"], args.get("columns"))
    elif name == "peek_values":
        result = peek_values(db_path, args["table"], args["column"], args.get("keyword"))
    elif name == "find_value":
        result = find_value(args["value"], bench, db_id)
    elif name == "read_skill":
        result = read_skill(SKILLS_DIR, args["name"])
    elif name == "ask_user":
        result = ask_user(args["question"], args.get("options") or [])
    elif name == "calculate":
        result = calculate(args["formula"], args.get("args") or {}, db_path=db_path)
    elif name == "chart":
        result = chart(args.get("spec") or {}, args.get("chart_type"),
                       args.get("x"), args.get("y"), db_path,
                       mcp_bridge=state.get("mcp_bridge"),
                       tenant=state.get("tenant", 1), _return_query=True)
        internal_query = result.pop("_query_result", None)
        if internal_query and internal_query.get("status") in ("ok", "empty"):
            trace.log("execute", turn=turn, tool="query_metrics",
                      sql=internal_query.get("sql"), n_rows=internal_query.get("n_rows", 0))
    elif name == "propose_metric_change":
        try:
            result = metric_changes.propose(
                os.path.dirname(db_path), args.get("metric"), args.get("field"),
                args.get("new_value"), args.get("reason"),
                state.get("session_id") or "agent")
        except metric_changes.MetricChangeError as exc:
            result = {"error": "metric_change_rejected", "reason": str(exc)}
    elif name == "save_report":
        result = save_report(args.get("title", ""), args.get("content", ""), state)
    elif name in mcp_names:
        # 中力度拦截: 已被语义层接管的平台工具, 裸调一律引导走 query_metrics
        reg = load_mdl(os.path.join(os.path.dirname(db_path), "mdl.yaml"), db_path=db_path)
        claimed = {m.get("mcp_tool") for m in reg.models.values() if m.get("mcp_tool")} \
            if reg else set()
        if name in claimed:
            result = {"error": "use_query_metrics",
                      "tool": name,
                      "reason": f"该工具的取数已由语义层接管(登记在 mdl.yaml: mcp_tool={name})。",
                      "how": ("请改用 query_metrics 点菜单查相应指标, 如: "
                              '{"metrics": ["访客数"], "time_range": {"start": "...", "end": "..."}}；'
                              "指标名见材料包'可用指标'。"),
                      "note": "口径/单位/维度由语义层保证; 直接调用会绕过这些保证, 故被拦截"}
        else:
            result = state["mcp_bridge"].call(name, args)
    elif name == "query_db":
        rendered = render_intent(args, db_path)
        if rendered.get("error"):
            result = rendered
        else:
            g = gate_check(rendered["sql"], db_path, bench, db_id)
            result = run_sql(rendered["sql"], db_path) if g.get("ok") else g
            if result.get("status") in ("ok", "empty"):
                result["sql"] = rendered["sql"]
                trace.log("execute", turn=turn, sql=rendered["sql"], n_rows=result.get("n_rows", 0))
    elif name == "execute_sql":
        g = gate_check(args["sql"], db_path, bench, db_id)
        result = run_sql(args["sql"], db_path) if g.get("ok") else g
        if result.get("status") in ("ok", "empty"):
            result["sql"] = args["sql"]
            trace.log("execute", turn=turn, sql=args["sql"], n_rows=result.get("n_rows", 0))
    elif name == "query_metrics":
        result = run_query_metrics(
            args, db_path, mcp_bridge=state.get("mcp_bridge"),
            tenant=state.get("tenant", 1))
        if result.get("status") in ("ok", "empty"):
            trace.log("execute", turn=turn, sql=result.get("sql"), n_rows=result.get("n_rows", 0))
    else:
        result = {"error": "unknown_tool", "available": ["peek_table", "peek_values", "query_db", "query_metrics", "execute_sql", "find_value", "read_skill", "ask_user", "calculate", "chart", "propose_metric_change", "save_report"] + list(mcp_names)}

    # chart 内部的 query_metrics 没有再绕一遍 dispatch，但仍必须产生同样的 Qn 账本记录。
    if internal_query and internal_query.get("status") in ("ok", "empty", "no_data"):
        state["q_count"] += 1
        ref = f"Q{state['q_count']}"
        internal_query["query_ref"] = ref
        result["query_ref"] = ref
        state["executions"][ref] = {
            "tool": "query_metrics", "args": args.get("spec") or {},
            "sql": internal_query.get("sql"),
            "n_rows": internal_query.get("n_rows"),
            "columns": internal_query.get("columns") or [],
            "rows": internal_query.get("rows") or [],
            "payload": {k: v for k, v in internal_query.items()
                        if k in ("status", "undefined_cells", "undefined_reasons",
                                 "undefined_note", "provenance")},
        }
    # 统一证据编号: 查询/计算结果进账本(存全量行, 供报告对账)
    # no_data 也是证据(数据缺失的 Qn 引用), 与 ok/empty 同级
    if result.get("status") in ("ok", "empty", "no_data") or name in mcp_names:
        state["q_count"] += 1
        ref = f"Q{state['q_count']}"
        result["query_ref"] = ref
        state["executions"][ref] = {"tool": name, "args": args,
                                    "sql": result.get("sql"),
                                    "n_rows": result.get("n_rows"),
                                    "columns": result.get("columns") or [],
                                    "rows": result.get("rows_all") or result.get("rows") or [],
                                    "payload": {k: v for k, v in result.items()
                                                if k in ("status", "summary", "funnel", "by_source",
                                                         "values", "daily", "value", "display",
                                                         "total_delta", "weight", "contributions",
                                                         "shares", "closure", "residual", "warnings",
                                                         "reconciliation")}}
    result.pop("rows_all", None)
    slim = {k: (v if k != "rows" else v[:3]) for k, v in result.items()}
    trace.log("tool_result", turn=turn, name=name, summary=slim)
    return result


NUMBER_LIKE = re.compile(r"(\d[\d,\.]{1,})|(\d+\s*(%|元|单|件|双|人))")
DATE_LIKE = re.compile(r"\d{4}[-/年]\d{1,2}([-/月]\d{1,2}日?)?|\d{1,2}[:：]\d{2}")


def validate_answer(answer: str, executions: dict) -> list:
    """答案证据标注的机械校验: Qn 引用必须真实存在; 有执行却无标注/无执行却给数字 → 打回."""
    text = DATE_LIKE.sub("", answer)  # 日期/时间不算数字结论
    issues = []
    refs = set(re.findall(r"\(Q(\d+)\)", answer))
    for ref in refs:
        if f"Q{ref}" not in executions:
            issues.append(f"答案引用了 (Q{ref}) 但该执行不存在")
    has_ref = bool(refs & {k[1:] for k in executions})
    has_number = bool(NUMBER_LIKE.search(text))
    if executions and not has_ref and has_number:
        issues.append("答案中的数字结论未标注依据编号(Qn), 请为每个数字补上引用")
    if not executions and has_number:
        issues.append("没有任何成功执行却给出了数字结论, 禁止编造")
    return issues


def resolve_db(bench: str, db_id: str) -> str:
    """兼容两种布局: <bench>/<db_id>.sqlite 或 <bench>/<db_id>/<db_id>.sqlite."""
    base = BENCH_DIRS[bench]
    for p in (os.path.join(base, f"{db_id}.sqlite"),
              os.path.join(base, db_id, f"{db_id}.sqlite")):
        if os.path.exists(p):
            return p
    return None


def append_audit(question: str, session: str, run_id: str, tenant: int,
                 executions: dict, status: str, refs=None) -> dict:
    """追加一条会话审计记录；审计失败不影响用户请求。"""
    refs = list(refs) if refs is not None else list(executions)
    current = [executions[ref] for ref in refs if ref in executions]
    tools_used = []
    for entry in current:
        tool = entry.get("tool")
        if tool and tool not in tools_used:
            tools_used.append(tool)
    record = {
        "ts": datetime.now().isoformat(timespec="milliseconds"),
        "session": session,
        "run_id": run_id,
        "question": question,
        "tenant": tenant,
        "tools_used": tools_used,
        "n_queries": len(current),
        "status": status,
    }
    try:
        os.makedirs("runs", exist_ok=True)
        with open(os.path.join("runs", "audit.jsonl"), "a", encoding="utf-8") as fh:
            fh.write(json.dumps(record, ensure_ascii=False, default=str) + "\n")
    except OSError:
        # 审计是旁路治理能力，不能把已完成的分析变成服务错误。
        pass
    return record


def run_agent(question: str, bench: str, db_id: str, budget: int = BUDGET,
              skills_dir: str = SKILLS_DIR, session_id: str = None,
              tenant: int = 1) -> dict:
    db_path = resolve_db(bench, db_id)
    if not db_path:
        return {"status": "error", "reason": f"库不存在: {bench}/{db_id}"}
    run_id = claim_dir("runs", "run")
    trace = Trace(os.path.join("runs", run_id, "trace.jsonl"))
    asked_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    trace.log("start", question=question, db_id=db_id, bench=bench)

    # 会话: 对话即上下文。已有会话加载历史; 新会话从材料包起头。
    sid = session_id or claim_dir("sessions", "s")
    sess_dir = os.path.join("sessions", sid)
    os.makedirs(sess_dir, exist_ok=True)
    sess_file = os.path.join(sess_dir, "messages.json")
    exec_file = os.path.join(sess_dir, "executions.json")
    sys_prompt = SYSTEM_PROMPT.format(
        database_context=f"当前数据库: SQLite, 库名 {db_id}。方言以 SQLite 为准。")
    state = {"q_count": 0, "executions": {}}
    if os.path.exists(exec_file):                    # 会话级连续编号: 第 2 问从上一问的 Qn 继续
        state = json.load(open(exec_file))
    state["session_id"] = sid
    state["db_path"] = db_path
    state["data_cutoff"] = get_data_cutoff(db_path)
    if os.path.exists(sess_file):
        messages = json.load(open(sess_file))
        messages.append({"role": "user",
                         "content": f"追问: {question}\n[提问时间: {asked_at}]"})
        resumed = True
    else:
        messages = [
            {"role": "system", "content": sys_prompt},
            {"role": "user",
             "content": f"问题: {question}\n\n{build_material(db_path, db_id, skills_dir)}"
                        + f"\n\n[提问时间: {asked_at}]"},
        ]
        resumed = False
    trace.log("session", sid=sid, resumed=resumed)
    client = OpenAI(api_key=os.environ["DEEPSEEK_API_KEY"], base_url="https://api.deepseek.com")

    # 外部平台数据: 通过 MCP 协议连接 mock 电商后台(独立进程, 物理隔离)
    state["mcp_tools"], state["mcp_bridge"] = {}, None
    extra_tools = []
    mcp_server = os.path.join("tools", "platform_mcp.py")
    if os.path.exists(mcp_server) and os.path.exists("data/platform/platform_data.sqlite"):
        bridge = McpBridge(sys.executable, ["tools/platform_mcp.py"], cwd=".")
        mcp_tools = bridge.start()
        state["mcp_bridge"] = bridge
        state["mcp_tools"] = {t["name"]: t for t in mcp_tools}
        extra_tools = [{"type": "function", "function": {
            "name": t["name"], "description": t["description"], "parameters": t["parameters"]}}
            for t in mcp_tools]
        trace.log("mcp_connect", tools=[t["name"] for t in mcp_tools])

    last_error_sig, error_repeat = None, 0
    validation_bounced = 0
    report_saved_this_run = False
    for turn in range(1, budget + 1):
        resp = client.chat.completions.create(
            model=MODEL, messages=messages, tools=tool_schema() + extra_tools, temperature=0)
        msg = resp.choices[0].message
        usage = getattr(resp, "usage", None)
        trace.log("llm", turn=turn,
                  usage={"p": usage.prompt_tokens, "c": usage.completion_tokens} if usage else None)
        assistant_msg = {"role": "assistant", "content": msg.content or ""}
        if msg.tool_calls:
            assistant_msg["tool_calls"] = [tc.model_dump() for tc in msg.tool_calls]
        messages.append(assistant_msg)

        if not msg.tool_calls:
            if last_error_sig:
                messages.append({"role": "user",
                                 "content": "上一步有工具错误未解决, 请调用工具处理后再给出最终回答。"})
                last_error_sig, error_repeat = None, 0
                continue
            issues = validate_answer(msg.content or "", state["executions"])
            content = msg.content or ""
            if "###EVIDENCE###" not in content:
                issues.append("最终回答缺少分隔符 ###EVIDENCE###——分隔符前给用户结论, 后给审计证据")
            else:
                user_part = content.split("###EVIDENCE###", 1)[0]
                if not report_saved_this_run:      # 报告类型: 允许小标题/粗体, 走独立校则
                    if "**" in user_part:
                        issues.append("结论区出现 markdown 粗体(**), 用户视图应为纯文字")
                    if re.search(r"本轮|上轮|首轮", user_part):
                        issues.append("结论区出现'本轮/上轮/首轮'式含糊引用——Qn 会话级连续, 直接写 (Qn)")
                else:
                    issues += [f"报告对账: {x}" for x in check_report(content, state["executions"])]
            if issues and validation_bounced < 3:
                validation_bounced += 1
                trace.log("answer_validation", turn=turn, issues=issues)
                messages.append({"role": "user", "content": "最终回答需要修正: " + "；".join(issues)
                                 + "。格式: 结论 → 换行 → ###EVIDENCE### → 证据明细(含 Qn 引用)。"})
                continue
            parts = content.split("###EVIDENCE###", 1)
            conclusion = parts[0].strip()
            evidence = parts[1].strip()
            trace.log("report", conclusion=conclusion, evidence=evidence,
                      executions=state["executions"])
            json.dump(messages, open(sess_file, "w"), ensure_ascii=False, default=str)
            json.dump(state, open(exec_file, "w"), ensure_ascii=False, default=str)
            with open(os.path.join(sess_dir, "runs.jsonl"), "a") as f:
                f.write(json.dumps({"ts": asked_at, "question": question, "run_id": run_id,
                                    "conclusion": conclusion, "status": "ok"},
                                   ensure_ascii=False) + "\n")
            return {"status": "ok", "answer": conclusion, "answer_full": msg.content,
                    "turns": turn, "run_id": run_id, "session_id": sid,
                    "executions": state["executions"]}

        for tc in msg.tool_calls:
            name = tc.function.name
            args = json.loads(tc.function.arguments)
            result = dispatch(name, args, db_path, bench, db_id, trace, turn, state)
            if name == "save_report" and result.get("status") == "saved":
                report_saved_this_run = True
            sig = f"{name}:{json.dumps(args, sort_keys=True, ensure_ascii=False)}"
            if result.get("error"):
                error_repeat = error_repeat + 1 if sig == last_error_sig else 1
                last_error_sig = sig
                if error_repeat >= 2:
                    result = dict(result, hint="同一动作已连续失败两次, 请换思路")
            else:
                last_error_sig, error_repeat = None, 0
            messages.append({"role": "tool", "tool_call_id": tc.id,
                             "content": json.dumps(result, ensure_ascii=False, default=str)})

    trace.log("end", status="budget_exhausted")
    json.dump(messages, open(sess_file, "w"), ensure_ascii=False, default=str)
    json.dump(state, open(exec_file, "w"), ensure_ascii=False, default=str)
    with open(os.path.join(sess_dir, "runs.jsonl"), "a") as f:
        f.write(json.dumps({"ts": asked_at, "question": question, "run_id": run_id,
                            "status": "budget_exhausted"}, ensure_ascii=False) + "\n")
    return {"status": "budget_exhausted", "turns": budget, "run_id": run_id, "session_id": sid}


def main():
    argv = sys.argv
    if "--question" in argv and "--db" in argv:
        question = argv[argv.index("--question") + 1]
        db_id = argv[argv.index("--db") + 1]
        bench = argv[argv.index("--bench") + 1] if "--bench" in argv else "bird"
        session = argv[argv.index("--session") + 1] if "--session" in argv else None
        out = run_agent(question, bench, db_id, session_id=session)
        print(json.dumps(out, ensure_ascii=False, default=str))
        return 0 if out["status"] == "ok" else 1
    print("用法: python3 agent_loop.py --question <问题> --db <库名> [--bench bird] [--session <会话id>]")
    return 2


if __name__ == "__main__":
    sys.exit(main())
